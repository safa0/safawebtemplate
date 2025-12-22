from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import sys
import os

# Theme Colors extracted from website config
COLOR_BG = RGBColor(0xE8, 0xDC, 0xC4)    # khaki-light
COLOR_TITLE = RGBColor(0x8B, 0x73, 0x55) # khaki-dark
COLOR_TEXT = RGBColor(0x33, 0x33, 0x33)  # Dark Grey for readability
COLOR_ACCENT = RGBColor(0xA8, 0x99, 0x68) # accent

def apply_slide_theme(slide):
    # Set background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

def create_reimagined_deck(src_path, dest_path):
    if not os.path.exists(src_path):
        print(f"Source file not found: {src_path}")
        return

    print(f"Reading from {src_path}")
    prs_src = Presentation(src_path)
    prs_dest = Presentation()
    
    # We will use a simple layout for all slides for now: Title + Content
    
    for i, slide_src in enumerate(prs_src.slides):
        print(f"Processing slide {i+1}...")
        
        # Decide layout
        # If it's the first slide, maybe use Title Slide layout (0)
        # However, standard python-pptx empty deck might differ.
        # Usually: 0=Title, 1=Title+Content
        layout_idx = 0 if i == 0 else 1
        if layout_idx >= len(prs_dest.slide_layouts):
            layout_idx = 0 # fallback
            
        slide_layout = prs_dest.slide_layouts[layout_idx]
        slide_dest = prs_dest.slides.add_slide(slide_layout)
        
        apply_slide_theme(slide_dest)
        
        # Extract all text shapes
        text_shapes = []
        for shape in slide_src.shapes:
            if shape.has_text_frame:
                if shape.text and shape.text.strip():
                    text_shapes.append(shape)
        
        # Sort by vertical position (top to bottom)
        text_shapes.sort(key=lambda s: s.top)
        
        if not text_shapes:
            continue
            
        # Assume first is title
        title_shape = text_shapes[0]
        
        # get destination title shape
        dest_title = slide_dest.shapes.title
        if dest_title:
            dest_title.text = title_shape.text
            # Style title
            for paragraph in dest_title.text_frame.paragraphs:
                paragraph.font.color.rgb = COLOR_TITLE
                paragraph.font.name = 'Arial'
                paragraph.font.bold = True
                if i == 0:
                    paragraph.font.size = Pt(44)
                else:
                    paragraph.font.size = Pt(32)
                    
        # Remaining shapes go into content placeholder or new text boxes
        body_shapes = text_shapes[1:]
        
        if body_shapes:
            content_placeholder = None
            # standard layouts usually have placeholders[1] as body
            # verify it's a body placeholder
            if len(slide_dest.placeholders) > 1:
                content_placeholder = slide_dest.placeholders[1]
            
            if content_placeholder and content_placeholder.has_text_frame:
                tf = content_placeholder.text_frame
                # We start fresh? Or clear it?
                # It's usually empty in a new slide
                tf.text = "" # ensure clear
                
                for src_shape in body_shapes:
                    # Append text from this shape
                    for src_p in src_shape.text_frame.paragraphs:
                        new_p = tf.add_paragraph()
                        new_p.text = src_p.text
                        new_p.level = src_p.level
                        new_p.font.color.rgb = COLOR_TEXT
                        new_p.font.name = 'Arial'
                        new_p.font.size = Pt(18)
                        
                        # Handle basic formatting if needed
                        # for run in src_p.runs: ... (complex, skip for now)
            else:
                # If no placeholder, create text boxes manualy
                top_offset = Inches(2)
                for src_shape in body_shapes:
                    left = Inches(1)
                    width = Inches(8)
                    height = Inches(1) 
                    
                    txBox = slide_dest.shapes.add_textbox(left, top_offset, width, height)
                    tf = txBox.text_frame
                    tf.word_wrap = True
                    tf.text = src_shape.text
                    
                    # Style
                    for p in tf.paragraphs:
                        p.font.color.rgb = COLOR_TEXT
                        p.font.name = 'Arial'
                        p.font.size = Pt(18)
                    
                    top_offset += Inches(1.5)

    prs_dest.save(dest_path)
    print(f"Saved to {dest_path}")

if __name__ == "__main__":
    create_reimagined_deck("Terneshk_PitchDeckMain.pptx", "Terneshk_Reimagined.pptx")
